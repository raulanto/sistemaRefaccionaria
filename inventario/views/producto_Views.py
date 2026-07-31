from django.http import JsonResponse
from django.urls import reverse_lazy
from django.views import View
from django.shortcuts import render, get_object_or_404, redirect
from inventario.forms.producto_form import ProductoForm
from inventario.models.movimientoInventario_modelo import MovimientoInventario
from inventario.models.productoImagen_modelo import ProductoImagen
from inventario.models.producto_modelo import Producto
from django.views.generic.edit import CreateView, UpdateView
from django.db.models import F, ProtectedError
import openpyxl
from openpyxl.styles import Font, PatternFill
from django.http import HttpResponse
from inventario.models import Producto
from django.http import JsonResponse
from django.db.models import Q
from inventario.services.inventario_service import InventarioService


class ProductoListaView(View):
    def get(self, request):
        productos = Producto.objects.all()
        data = {
            "titulo": "Sistema de Inventario",
            "mensaje": "Bienvenido al sistema de inventario de refacciones.",
            "productos": productos,
        }
        return render(request, "producto/producto_lista.html", context=data)


class ProductoDetalleView(View):
    def get(self, request, id_producto):
        producto = get_object_or_404(Producto, id=id_producto)
        imagenes = ProductoImagen.objects.filter(producto=producto).order_by("orden")
        data = {
            "titulo": "Detalle del Producto",
            "producto": producto,
            "imagenes": imagenes,
        }
        return render(request, "producto/producto_detalle.html", context=data)


class ProductoCrearView(CreateView):
    model = Producto
    form_class = ProductoForm
    template_name = "producto/producto_crear.html"
    success_url = reverse_lazy("inventario:productoLista")
    
    def get_context_data(self, **kwargs):
            context = super().get_context_data(**kwargs)
            context["titulo"] = "Crear Producto"
            context["mensaje"] = "Completa el formulario para crear un nuevo producto"
            return context

    def form_valid(self, form):
        response = super().form_valid(
            form
        )  # el producto ya se guardó con su stock inicial

        stock_inicial = self.object.stock
        if stock_inicial > 0:
            # Solo registramos el movimiento para que quede en el Kardex,
            # SIN volver a sumar al stock (ya está guardado desde el formulario)
            MovimientoInventario.objects.create(
                producto=self.object,
                tipo="ENTRADA",
                cantidad=stock_inicial,
                motivo="AJUSTE",
                observaciones="Stock inicial al dar de alta el producto",
                usuario=(
                    self.request.user if self.request.user.is_authenticated else None
                ),
            )
        return response


class ProductoEditarView(UpdateView):
    model = Producto
    form_class = ProductoForm
    template_name = "producto/producto_crear.html"
    success_url = "/producto/producto_lista/"


def EliminarProductoAjax(request, pk):
    if request.method == "POST":
        producto = get_object_or_404(Producto, pk=pk)
        try:
            producto.delete()
            return JsonResponse({"estado": "ok"})
        except ProtectedError:
            return JsonResponse(
                {
                    "estado": "error",
                    "mensaje": "No se puede eliminar porque el producto ya fue vendido.",
                }
            )


def ExportarProdutosExcel(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Productos"

    # Encabezados
    encabezados = [
        "Código", "Nombre", "Categoría", "Marca", "Proveedor",
        "Precio Compra", "Precio Venta", "Stock", "Stock Mínimo", "Estado"
    ]
    ws.append(encabezados)

    # Estilo del encabezado
    for col_num, _ in enumerate(encabezados, 1):
        celda = ws.cell(row=1, column=col_num)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")

    # Datos
    productos = Producto.objects.select_related(
        "categoria", "marca", "proveedor_principal"
    ).all()

    for p in productos:
        ws.append([
            p.codigo,
            p.nombre,
            p.categoria.nombre,
            p.marca.nombre,
            p.proveedor_principal.nombre,
            float(p.precio_compra),
            float(p.precio_venta),
            p.stock,
            p.stock_minimo,
            p.estado,
        ])

    # Ajustar ancho de columnas automáticamente
    for col in ws.columns:
        max_len = max(len(str(c.value)) for c in col if c.value is not None)
        ws.column_dimensions[col[0].column_letter].width = max_len + 3

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="productos.xlsx"'
    wb.save(response)
    return response

def buscar_productos_ajax(request):
    """
    Busca productos por código o nombre para los buscadores del sistema.
    GET /productos/buscar_ajax/?q=balata
    """
    q = request.GET.get('q', '').strip()

    productos = Producto.objects.all()

    if q:
        productos = productos.filter(
            Q(codigo__icontains=q) | Q(nombre__icontains=q)
        )

    productos = productos.order_by('nombre')[:20]

    data = [
        {
            'id': p.id,
            'codigo': p.codigo,
            'nombre': p.nombre,
            'precio': float(p.precio_venta),
            'stock': p.stock,
            'tiene_iva': p.tiene_iva,
        }
        for p in productos
    ]
    return JsonResponse({'productos': data})
